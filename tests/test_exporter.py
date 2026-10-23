from docx import Document

from marketresearch.report.exporter import markdown_to_docx

MD = """# Title Here

> **DEMO MODE.** banner

## 1. Section

- bullet with **bold** and *italic* [1]
- second bullet

| A | B |
|---|---|
| x | y |
| z | w |

1. numbered item

Plain paragraph.
"""


def test_markdown_to_docx_structure(tmp_path):
    out = markdown_to_docx(MD, tmp_path / "t.docx")
    d = Document(str(out))
    texts = [p.text for p in d.paragraphs]
    assert "Title Here" in texts and "1. Section" in texts
    assert len(d.tables) == 1 and d.tables[0].cell(0, 0).text == "A" and d.tables[0].cell(2, 1).text == "w"
    assert any("bullet with bold and italic [1]" == t for t in texts)
