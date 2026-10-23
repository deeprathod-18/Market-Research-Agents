"""Minimal, dependency-light Markdown -> DOCX converter tailored to the brief's structure."""
from __future__ import annotations

import re
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x4E, 0x79)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _runs(par, text: str, size: float | None = None, color: RGBColor | None = None, bold_all: bool = False) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        bold = italic = code = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        r = par.add_run(part)
        r.bold = bold or bold_all
        r.italic = italic
        if code:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def _add_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def _table(doc: Document, lines: List[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines if not re.match(r"^\|\s*:?-{2,}", ln)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            _runs(par, row[j] if j < len(row) else "", size=8.5, color=RGBColor(255, 255, 255) if i == 0 else None, bold_all=(i == 0))
            if i == 0:
                _shade(cell, "1F4E79")
            elif i % 2 == 0:
                _shade(cell, "EAF1F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def markdown_to_docx(md: str, out_path: Path, base_dir: Path | None = None) -> Path:
    base_dir = Path(base_dir or out_path.parent)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    _add_page_number(sec)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(size), NAVY, True

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and not any(p.style.name == "Title" for p in doc.paragraphs):
                p = doc.add_paragraph(style="Title")
                p.add_run(text).font.color.rgb = NAVY
            else:
                doc.add_heading(text, level=min(level, 3))
            i += 1
            continue
        img = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", ln.strip())
        if img:
            path = Path(img.group(2))
            path = path if path.is_absolute() else base_dir / path
            if path.exists():
                doc.add_picture(str(path), width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _runs(cap, f"*Figure: {img.group(1)}*", size=8.5)
            i += 1
            continue
        if ln.startswith(">"):
            p = doc.add_paragraph()
            _runs(p, ln.lstrip("> ").strip(), size=10)
            p.paragraph_format.left_indent = Inches(0.2)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "FFF2CC")
            pPr.append(shd)
            i += 1
            continue
        b = re.match(r"^\s*[-*]\s+(.*)", ln)
        if b:
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, b.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        n = re.match(r"^\s*\d+[.)]\s+(.*)", ln)
        if n:
            p = doc.add_paragraph(style="List Number")
            _runs(p, n.group(1))
            i += 1
            continue
        p = doc.add_paragraph()
        _runs(p, ln.strip())
        p.paragraph_format.space_after = Pt(4)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
